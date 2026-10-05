"""Regenerate the binary test fixtures. Run from backend/:  uv run python -m tests.fixtures.build_fixtures

- resumes/*.pdf, resumes/*.docx: fictional input resumes in deliberately varied (often ATS-unfriendly) styles.
- ats/*.pdf: one valid Sherlock-rendered resume plus deliberately broken variants for ats_check tests,
  and ats/expectations.json describing what the valid render contains.

All people, emails and phone numbers here are fictional.
"""

import asyncio
import json
import subprocess
import tempfile
from pathlib import Path

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.text.paragraph import Paragraph

from app.schemas.resume import MasterResume
from app.services.resume.render import FONTS_DIR, compile_typst, render_resume

FIXTURES = Path(__file__).parent
RESUMES = FIXTURES / "resumes"
ATS = FIXTURES / "ats"


def typst(src: Path, out: Path, *extra: str) -> None:
    subprocess.run(
        [
            "typst",
            "compile",
            "--ignore-system-fonts",
            "--font-path",
            str(FONTS_DIR),
            *extra,
            str(src),
            str(out),
        ],
        check=True,
    )


# --- DOCX input resume ---------------------------------------------------------------------------


def _hyperlink(paragraph: Paragraph, url: str, text: str) -> None:
    rel_id = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    text_el = OxmlElement("w:t")
    text_el.text = text
    run.append(text_el)
    link.append(run)
    paragraph._p.append(link)


def build_docx(path: Path) -> None:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10)

    title = doc.add_paragraph()
    title.add_run("Rohan Das").bold = True
    links = doc.add_paragraph()
    for i, (url, text) in enumerate(
        [
            ("https://linkedin.com/in/rohandas-oss", "linkedin.com/in/rohandas-oss"),
            ("https://github.com/rohan-das-oss", "github.com/rohan-das-oss"),
            ("https://rohandas.example.dev", "rohandas.example.dev"),
        ]
    ):
        if i:
            links.add_run(" | ")
        _hyperlink(links, url, text)
    links.add_run(" | rohan.das@example.com | +91 90000 33333")

    def heading(text: str) -> None:
        doc.add_paragraph().add_run(text).bold = True

    def entry(left: str, right: str) -> None:
        p = doc.add_paragraph()
        p.add_run(left).bold = True
        p.add_run(f"\t{right}")

    def bullet(text: str) -> None:
        doc.add_paragraph(text, style="List Bullet")

    heading("EDUCATION")
    entry("Bachelor of Technology, Computer Science and Engineering", "2025 - 2029")
    doc.add_paragraph("Example School of Technology, Bengaluru\tCGPA: 8.61/10")

    heading("EXPERIENCE")
    entry("Open Source Contributor - Example Office Suite", "Jan 2026 - Present")
    bullet(
        "Fixed a long-standing spreadsheet bug where auto-filled decimal series produced rounding errors "
        "(0.30000000000000004 instead of 0.3); added a regression test, merged after five review rounds"
    )
    bullet(
        "Diagnosed a rendering bug shared by the word processor and presentation apps by tracing a shared layout module"
    )
    entry("Open Source Contributor - Example Observability Project", "May 2026 - Present")
    bullet(
        "Built a registry crawler for JavaScript instrumentation packages, tracking 52 entries, "
        "merged after 12 rounds of maintainer review"
    )
    bullet("Fixed a telemetry SDK bug that sent local development data to production, causing CORS errors")
    bullet("Merged 18+ pull requests and wrote onboarding documentation for new contributors")

    heading("PROJECTS")
    entry("TxScope - Bitcoin Transaction Explorer", "2026")
    doc.add_paragraph("github.com/rohan-das-oss/txscope")
    bullet("Wrote a raw Bitcoin transaction parser with no external libraries, supporting SegWit and Taproot")
    bullet("Exposed the parser as a CLI, a REST API and a web visualiser")
    entry("Nyaya Search - Retrieval System for Legal Documents", "2026")
    doc.add_paragraph("github.com/rohan-das-oss/nyaya-search")
    bullet("Combined BM25 keyword search with embedding search and re-ranking so every answer cites a source")
    bullet("Built a 90-question evaluation suite run in CI that blocks merges which reduce accuracy")

    heading("SKILLS")
    doc.add_paragraph("Languages: C++, Python, Java, JavaScript, TypeScript")
    doc.add_paragraph("Backend / Frontend: Node.js, Express.js, Django, React, Next.js, HTML, CSS")
    doc.add_paragraph("Databases: PostgreSQL, MongoDB    Tools: Git, Docker, GitHub Actions, Postman")

    heading("EXTRA-CURRICULAR ACTIVITIES")
    entry("Student Developer Ambassador, Example School of Technology", "May 2026 - Jul 2026")
    bullet("Organised workshops and tech talks on cloud developer tools for 200+ students")
    doc.save(str(path))


# --- ATS fixtures ----------------------------------------------------------------------------------

BASE_RESUME = MasterResume.model_validate(
    {
        "basics": {
            "name": "Kavya Raman",
            "email": "kavya.raman@example.com",
            "phone": "+91 90000 44444",
            "location": "Bengaluru, India",
            "links": [{"label": "GitHub", "url": "https://github.com/kavyaraman"}],
        },
        "summary": "Backend-focused CS student who builds efficient, well-tested services and workflow tools.",
        "education": [
            {
                "id": "edu1",
                "institution": "Example Institute of Technology",
                "degree": "B.Tech",
                "field": "Computer Science",
                "start": "2023-08",
                "end": "2027-05",
                "gpa": "8.7/10",
                "bullets": [],
            }
        ],
        "experience": [
            {
                "id": "exp1",
                "org": "Example Fintech",
                "role": "Software Engineering Intern",
                "location": "Bengaluru",
                "start": "2026-05",
                "end": "2026-07",
                "bullets": [
                    {
                        "id": "exp1-b1",
                        "text": "Built an efficient reconciliation workflow in Python and PostgreSQL that "
                        "cut manual review time by 35% for the finance team",
                        "skills": ["python", "postgresql"],
                    },
                    {
                        "id": "exp1-b2",
                        "text": "Added profiling and fixed slow queries in the payments service, reducing "
                        "p95 latency from 420 ms to 180 ms",
                        "skills": ["postgresql"],
                    },
                ],
            }
        ],
        "projects": [
            {
                "id": "proj1",
                "name": "Flowlog",
                "link": "github.com/kavyaraman/flowlog",
                "tech": ["Go", "Redis"],
                "start": "2025-11",
                "bullets": [
                    {
                        "id": "proj1-b1",
                        "text": "Designed a log shipper in Go with Redis-backed buffering that handles "
                        "fluctuating traffic of 5k lines per second",
                        "skills": ["go", "redis"],
                    },
                    {
                        "id": "proj1-b2",
                        "text": "Wrote a configuration file linter that flags conflicting filters before deploys",
                        "skills": ["go"],
                    },
                ],
            }
        ],
        "skills": {"Languages": ["Python", "Go", "SQL"], "Tools": ["PostgreSQL", "Redis", "Docker", "Git"]},
        "achievements": [
            {"id": "ach-b1", "text": "Finalist, national student hackathon 2025 (top 10 of 400 teams)"},
        ],
    }
)


async def build_ats_fixtures() -> None:
    ATS.mkdir(exist_ok=True)
    rendered = await render_resume(BASE_RESUME)
    (ATS / "valid.pdf").write_bytes(rendered.pdf)
    source = rendered.typst_source

    # Ligatures: Typst maps ligature glyphs back to "fi"/"fl" on extraction, so merely enabling them
    # doesn't reproduce the problem. Tools like pdfLaTeX (without cmap) and some Word exports instead emit
    # presentation-form code points (U+FB01 "\ufb01"...), which break keyword matching in ATS parsers.
    # Reproduce that: ligatures on, and the ligature code points written into the text itself.
    preamble_end = source.index("#block(below")
    body = source[preamble_end:]
    for plain, glyph in (("ffi", "\ufb03"), ("ffl", "\ufb04"), ("fi", "\ufb01"), ("fl", "\ufb02")):
        body = body.replace(plain, glyph)
    ligatures = source[:preamble_end].replace("ligatures: false,", "ligatures: true,", 1) + body
    assert "\ufb01" in ligatures
    (ATS / "ligatures.pdf").write_bytes(await compile_typst(ligatures))

    # Two columns: a sidebar (education, skills, achievements) beside the main column (experience,
    # projects), the classic "designed" resume layout. Text from both columns shares the same lines.
    def section(name: str) -> str:
        start = source.index(f'= #"{name}"')
        nxt = source.find("\n= #", start + 1)
        return source[start : nxt + 1 if nxt != -1 else len(source)]

    head = source[: source.index("= #")]
    sidebar = "".join(section(n) for n in ("Education", "Skills", "Achievements"))
    main = "".join(section(n) for n in ("Experience", "Projects"))
    two_column = (
        head + f"#grid(columns: (1fr, 1.6fr), column-gutter: 1.5em,\n[\n{sidebar}],\n[\n{main}],\n)\n"
    )
    (ATS / "two_column.pdf").write_bytes(await compile_typst(two_column))

    # Contact details only in the page header (ATS often skip headers).
    block_start = source.index("#block(below: 0.6em)[")
    contact_block = source[block_start : source.index("\n]\n", block_start) + 3]
    header = contact_block.replace("#block(below: 0.6em)[", "[", 1).strip()
    header_source = source.replace(contact_block, "").replace(
        '#set page(paper: "a4", margin: (x: 0.6in, y: 0.5in))',
        f'#set page(paper: "a4", margin: (x: 0.6in, top: 1.1in, bottom: 0.5in), header: {header})',
        1,
    )
    (ATS / "header_contact.pdf").write_bytes(await compile_typst(header_source))

    # Image-only "scan": the valid page rasterised and placed as a picture, no text layer.
    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / "valid.typ"
        src.write_text(source)
        typst(src, Path(tmp) / "page.png", "--format", "png", "--ppi", "110")
        scan = Path(tmp) / "scan.typ"
        scan.write_text('#set page(paper: "a4", margin: 0pt)\n#image("page.png", width: 100%)\n')
        typst(scan, ATS / "image_only.pdf")

    expectations = {
        "name": rendered.name,
        "email": rendered.email,
        "phone": rendered.phone,
        "headings": rendered.headings,
        "bullets": rendered.bullets,
    }
    (ATS / "expectations.json").write_text(json.dumps(expectations, indent=2) + "\n")


def main() -> None:
    RESUMES.mkdir(exist_ok=True)
    for name in ("aarav_menon", "meera_iyer"):
        typst(RESUMES / "src" / f"{name}.typ", RESUMES / f"{name}.pdf")
    build_docx(RESUMES / "rohan_das.docx")
    asyncio.run(build_ats_fixtures())
    print("fixtures rebuilt")


if __name__ == "__main__":
    main()
