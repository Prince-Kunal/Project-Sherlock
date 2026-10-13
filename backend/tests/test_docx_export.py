"""DOCX export (PLAN.md Phase 10): the Word version of a resume passes an equivalent ATS check, and a
file an ATS would misread is refused."""

import io
import uuid

from docx import Document
from httpx import AsyncClient
from sqlalchemy.ext.asyncio import AsyncSession

from app.models import User
from app.services.resume.docx_export import DOCX_MEDIA_TYPE, checked_docx, docx_check, render_docx
from app.services.resume.render import build_context
from tests.matching_fixtures import fixture_resume
from tests.sending_fixtures import sendable


def test_the_word_resume_mirrors_the_pdf_content_and_passes() -> None:
    resume = fixture_resume()
    order = ["experience", "projects", "skills", "education"]
    rendered, report = checked_docx(resume, order)
    assert report.passed, report.failures
    assert rendered.filename == "Rohan_Das_Resume.docx"
    assert rendered.headings == ["Experience", "Projects", "Skills", "Education"]
    expected = [
        b
        for s in build_context(resume, order)["sections"]
        for e in s.get("entries", [])
        for b in e["bullets"]
    ]
    assert sorted(rendered.bullets) == sorted(expected)

    document = Document(io.BytesIO(rendered.docx))
    assert document.core_properties.title == "Rohan Das — Resume"
    assert document.core_properties.author == "Rohan Das"
    assert document.tables == []
    text = "\n".join(p.text for p in document.paragraphs)
    assert text.splitlines()[0] == "Rohan Das"
    assert "rohan@example.com" in "\n".join(text.splitlines()[:3])
    assert all(section.header.paragraphs[0].text == "" for section in document.sections)


def _rewrite(data: bytes, change: object) -> bytes:
    document = Document(io.BytesIO(data))
    change(document)  # type: ignore[operator]
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def test_the_check_refuses_what_an_ats_would_misread() -> None:
    rendered = render_docx(fixture_resume())
    expect = rendered.expectations()

    def table(document: Document) -> None:  # type: ignore[valid-type]
        document.add_table(rows=1, cols=2)

    def header(document: Document) -> None:  # type: ignore[valid-type]
        document.sections[0].header.paragraphs[0].text = "Rohan Das, rohan@example.com"

    def ligature(document: Document) -> None:  # type: ignore[valid-type]
        bullet = next(p for p in document.paragraphs if p.text.startswith("•"))
        bullet.runs[0].text = bullet.runs[0].text.replace("fi", "ﬁ", 1).replace("e", "e", 1) + " ﬁ"

    def contact_last(document: Document) -> None:  # type: ignore[valid-type]
        contact = document.paragraphs[1]
        contact._p.getparent().remove(contact._p)
        document.add_paragraph("rohan@example.com")

    failures = {
        name: docx_check(_rewrite(rendered.docx, change), expect).failures
        for name, change in [
            ("table", table),
            ("header", header),
            ("ligature", ligature),
            ("contact", contact_last),
        ]
    }
    assert "DOCX contains a table" in failures["table"]
    assert any("header or footer" in f for f in failures["header"])
    assert any("ligature" in f for f in failures["ligature"])
    assert any("email not found in the first 15%" in f for f in failures["contact"])


async def test_downloads(auth_client: AsyncClient, session: AsyncSession) -> None:
    from tests.test_matches_api import scored_user

    me = uuid.UUID((await auth_client.get("/auth/me")).json()["id"])
    assert (await auth_client.get("/resume/preview.docx")).status_code == 404
    await scored_user(session, me)
    master = await auth_client.get("/resume/preview.docx")
    assert master.status_code == 200
    assert master.headers["content-type"] == DOCX_MEDIA_TYPE
    assert 'filename="Rohan_Das_Resume.docx"' in master.headers["content-disposition"]

    user = await session.get(User, me)
    assert user is not None
    outreach = await sendable(session, user, approve=False)
    assert (await auth_client.get(f"/outreach/{outreach.id}/resume.docx")).status_code == 404  # no record
    outreach.tailored_resume = {
        **(outreach.tailored_resume or {}),
        "resume": fixture_resume().model_dump(mode="json"),
        "section_order": ["projects", "experience", "skills", "education"],
    }
    await session.commit()
    tailored = await auth_client.get(f"/outreach/{outreach.id}/resume.docx")
    assert tailored.status_code == 200
    lines = [p.text for p in Document(io.BytesIO(tailored.content)).paragraphs]
    assert lines.index("Projects") < lines.index("Experience")
