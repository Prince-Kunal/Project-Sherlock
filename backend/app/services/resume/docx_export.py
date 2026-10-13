"""DOCX export of a resume (PLAN.md Phase 10): some older ATS parse Word better than PDF.

Built from the same prepared content as the PDF (`render.build_context`: same sections, order,
bullets and wording), following §6.1: one column, no tables, text boxes, images or header/footer
content, contact details in the body, standard headings, "•" bullets typed as text (no Word list
numbering), links written out, Arial 10.5 pt, 0.6 in margins, and the document title/author set.

`docx_check` is the DOCX counterpart of the PDF ATS check: the structure rules above, then the text
rules (no unmappable or ligature characters, contact details in the first 15%, headings as their own
lines in order, every bullet intact) on two independent extractions: python-docx's paragraph text
and a raw walk of word/document.xml the way docx2txt-style parsers read it. Page count can't be
measured without a layout engine, so the one-page fit comes from the PDF it mirrors.
"""

import asyncio
import io
import re
import zipfile
from collections.abc import Sequence
from dataclasses import dataclass

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph
from fastapi import HTTPException, Response, status

from app.schemas.resume import MasterResume
from app.services.resume.ats_check import AtsExpectations, AtsReport, _check_text, keyword_coverage
from app.services.resume.render import build_context, resume_filename

FONT = "Arial"
BODY_PT = 10.5
NAME_PT = 18
BULLET = "• "
CONTACT_SEP = " | "
DOCX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


@dataclass
class RenderedDocx:
    docx: bytes
    filename: str
    name: str
    email: str | None
    phone: str | None
    headings: list[str]
    bullets: list[str]

    def expectations(self, keywords: Sequence[str] = ()) -> AtsExpectations:
        return AtsExpectations(
            name=self.name,
            email=self.email,
            phone=self.phone,
            headings=self.headings,
            bullets=self.bullets,
            keywords=list(keywords),
        )


def _add_hyperlink(paragraph: Paragraph, text: str, url: str) -> None:
    """A clickable link whose visible text is the URL itself (§6.1)."""
    part = paragraph.part
    rel_id = part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    text_el = OxmlElement("w:t")
    text_el.text = text
    text_el.set(qn("xml:space"), "preserve")
    run.append(text_el)
    link.append(run)
    paragraph._p.append(link)


def _para(
    document: object, text: str = "", *, bold: bool = False, size: float = BODY_PT, after: float = 0
) -> Paragraph:
    paragraph: Paragraph = document.add_paragraph()  # type: ignore[attr-defined]
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.space_before = Pt(0)
    if text:
        run = paragraph.add_run(text)
        run.bold = bold
        run.font.size = Pt(size)
    return paragraph


def render_docx(resume: MasterResume, section_order: Sequence[str] | None = None) -> RenderedDocx:
    context = build_context(resume, section_order)
    basics = resume.basics
    document = Document()
    style = document.styles["Normal"]
    style.font.name = FONT
    style.font.size = Pt(BODY_PT)
    for section in document.sections:
        section.left_margin = section.right_margin = Inches(0.6)
        section.top_margin = section.bottom_margin = Inches(0.6)
    document.core_properties.title = context["doc_title"]
    document.core_properties.author = basics.name

    name = _para(document, basics.name, bold=True, size=NAME_PT, after=2)
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact_line = _para(document, after=6)
    contact_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for i, item in enumerate(context["contact"]):
        if i:
            contact_line.add_run(CONTACT_SEP)
        if item["url"] and not str(item["url"]).startswith("mailto:"):
            _add_hyperlink(contact_line, item["text"], item["url"])
        else:
            contact_line.add_run(item["text"])

    if context["summary"]:
        _para(document, context["summary"], after=6)

    headings: list[str] = []
    bullets: list[str] = []
    for section in context["sections"]:
        headings.append(section["heading"])
        heading = _para(document, section["heading"], bold=True, size=BODY_PT + 1.5, after=2)
        heading.paragraph_format.space_before = Pt(6)
        if section["kind"] == "skills":
            for line in section["lines"]:
                paragraph = _para(document)
                paragraph.add_run(f"{line['category']}: ").bold = True
                paragraph.add_run(line["joined"])
            continue
        for entry in section["entries"]:
            if entry["title"]:
                _para(document, entry["title"], bold=True, after=0)
            if entry["meta"]:
                _para(document, entry["meta"], after=1)
            for text in entry["bullets"]:
                paragraph = _para(document, BULLET + text, after=1)
                paragraph.paragraph_format.left_indent = Inches(0.15)
                paragraph.paragraph_format.first_line_indent = Inches(-0.15)
                bullets.append(text)

    buffer = io.BytesIO()
    document.save(buffer)
    return RenderedDocx(
        docx=buffer.getvalue(),
        filename=resume_filename(basics.name).removesuffix(".pdf") + ".docx",
        name=basics.name,
        email=basics.email,
        phone=basics.phone,
        headings=headings,
        bullets=bullets,
    )


# --- The check -----------------------------------------------------------------------------------------


def _python_docx_text(data: bytes) -> str:
    document = Document(io.BytesIO(data))
    return "\n".join(p.text for p in document.paragraphs)


def _raw_xml_text(data: bytes) -> str:
    """word/document.xml read the way simple parsers do: every <w:t> in order, a newline per paragraph,
    tabs and breaks kept (independent of python-docx's own paragraph text)."""
    body = Document(io.BytesIO(data)).element.body
    lines = []
    for paragraph in body.iter(qn("w:p")):
        parts = []
        for node in paragraph.iter():
            if node.tag == qn("w:t"):
                parts.append(node.text or "")
            elif node.tag == qn("w:tab"):
                parts.append("\t")
            elif node.tag in (qn("w:br"), qn("w:cr")):
                parts.append("\n")
        lines.append("".join(parts))
    return "\n".join(lines)


def _structure_failures(data: bytes) -> list[str]:
    failures = []
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        names = archive.namelist()
        body = archive.read("word/document.xml").decode("utf-8", errors="replace")
        if any(n.startswith("word/media/") for n in names):
            failures.append("DOCX contains images")
        for part in (n for n in names if re.match(r"word/(header|footer)\d*\.xml$", n)):
            text = re.sub(r"<[^>]+>", "", archive.read(part).decode("utf-8", errors="replace")).strip()
            if text:
                failures.append(f"DOCX has text in a page header or footer ({part})")
    if "<w:tbl>" in body or "<w:tbl " in body:
        failures.append("DOCX contains a table")
    if "txbxContent" in body:
        failures.append("DOCX contains a text box")
    if "<w:drawing" in body or "<w:pict" in body:
        failures.append("DOCX contains a drawing")
    if "<w:numPr" in body:
        failures.append("DOCX uses Word list numbering for bullets")
    return failures


def docx_check(data: bytes, expect: AtsExpectations) -> AtsReport:
    failures = _structure_failures(data)
    plain = _python_docx_text(data)
    if not plain.strip():
        return AtsReport(passed=False, failures=[*failures, "DOCX has no text"], page_count=1)
    failures += _check_text("python-docx", plain, expect)
    failures += _check_text("document.xml", _raw_xml_text(data), expect)
    coverage = keyword_coverage(plain, expect.keywords) if expect.keywords else None
    return AtsReport(passed=not failures, failures=failures, page_count=1, keyword_coverage=coverage)


def checked_docx(
    resume: MasterResume, section_order: Sequence[str] | None = None
) -> tuple[RenderedDocx, AtsReport]:
    rendered = render_docx(resume, section_order)
    return rendered, docx_check(rendered.docx, rendered.expectations())


async def docx_response(resume: MasterResume, section_order: Sequence[str] | None) -> Response:
    """A download of the checked DOCX; `409 docx_check_failed` (with the reasons) if it fails the check,
    so a file an ATS would misread is never handed out."""
    rendered, report = await asyncio.to_thread(checked_docx, resume, section_order)
    if not report.passed:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail={
                "code": "docx_check_failed",
                "message": "The Word version didn't pass the ATS check: " + "; ".join(report.failures[:3]),
            },
        )
    return Response(
        content=rendered.docx,
        media_type=DOCX_MEDIA_TYPE,
        headers={
            "Content-Disposition": f'attachment; filename="{rendered.filename}"',
            "Cache-Control": "no-store",
        },
    )
