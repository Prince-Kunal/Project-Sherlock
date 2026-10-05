"""Uploaded resume (PDF/DOCX) → MasterResume (PLAN.md Phase 1).

Privacy (§3.1): email and phone are found deterministically and redacted before the text reaches the
LLM; they're put back into `basics` afterwards. Profile links are also found deterministically, from
both the visible text and hidden hyperlink targets (resumes often hide URLs behind words like "GitHub").
"""

import asyncio
import io
import re
import zipfile
from dataclasses import dataclass
from pathlib import PurePath
from urllib.parse import urlparse

import pdfplumber
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE

from app.schemas.resume import (
    Basics,
    Bullet,
    Education,
    Experience,
    Link,
    MasterResume,
    ParsedBullet,
    ParsedResume,
    Project,
)
from app.services.llm.client import LLMClient
from app.services.llm.prompt_loader import load_prompt
from app.services.llm.redaction import redact_pii
from app.services.resume.render import display_url
from app.services.resume.skills import SkillLexicon, get_lexicon

MAX_TEXT_CHARS = 20_000

_EMAIL_RE = re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")
_PHONE_RE = re.compile(r"(?<![\w])\+?\(?\d[\d\s().-]{8,}\d(?![\w])")
_URL_RE = re.compile(r"(?:https?://)?(?:www\.)?(?:[a-z0-9-]+\.)+[a-z]{2,}(?:/[^\s|,;)\]]*)?", re.IGNORECASE)
_PROFILE_PATTERNS: list[tuple[str, re.Pattern[str]]] = [
    ("LinkedIn", re.compile(r"^linkedin\.com/in/[^/]+$")),
    ("GitHub", re.compile(r"^github\.com/[^/]+$")),
    ("GitLab", re.compile(r"^gitlab\.com/[^/]+$")),
    ("LeetCode", re.compile(r"^leetcode\.com/(u/)?[^/]+$")),
    ("Kaggle", re.compile(r"^kaggle\.com/[^/]+$")),
    ("Codeforces", re.compile(r"^codeforces\.com/profile/[^/]+$")),
    ("HackerRank", re.compile(r"^hackerrank\.com/(profile/)?[^/]+$")),
]


_COMMON_TLDS = {
    "com",
    "org",
    "net",
    "io",
    "dev",
    "me",
    "tech",
    "ai",
    "app",
    "in",
    "co",
    "xyz",
    "info",
    "edu",
    "site",
    "online",
    "page",
    "codes",
    "blog",
    "live",
    "so",
    "sh",
    "gg",
    "cc",
    "us",
    "uk",
}


def _looks_like_url(candidate: str) -> bool:
    """Filter regex hits from free text: "Node.js" and "B.Tech" are not links."""
    if re.match(r"^(https?://|www\.)", candidate, re.IGNORECASE):
        return True
    host = urlparse(f"https://{candidate}").hostname or ""
    labels = host.lower().split(".")
    return len(labels) >= 2 and labels[-1] in _COMMON_TLDS and len(labels[-2]) >= 3


class ResumeExtractionError(Exception):
    """The file couldn't be read as a resume (wrong type, corrupt, or no text layer)."""


@dataclass
class ExtractedDocument:
    text: str
    link_targets: list[str]  # hyperlink URIs embedded in the file


def _extract_pdf(data: bytes) -> ExtractedDocument:
    try:
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            texts, links = [], []
            for page in pdf.pages:
                texts.append(page.extract_text() or "")
                links += [h["uri"] for h in page.hyperlinks if h.get("uri")]
    except Exception as exc:  # pdfminer raises a zoo of exception types for bad files
        raise ResumeExtractionError("could not read this PDF") from exc
    return ExtractedDocument(text="\n".join(texts), link_targets=links)


def _extract_docx(data: bytes) -> ExtractedDocument:
    try:
        doc = Document(io.BytesIO(data))
    except (zipfile.BadZipFile, KeyError, ValueError) as exc:
        raise ResumeExtractionError("could not read this Word file") from exc
    lines = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text for cell in row.cells))
    links = [rel.target_ref for rel in doc.part.rels.values() if rel.reltype == RELATIONSHIP_TYPE.HYPERLINK]
    return ExtractedDocument(text="\n".join(lines), link_targets=links)


def extract_document(data: bytes, filename: str) -> ExtractedDocument:
    suffix = PurePath(filename).suffix.lower()
    if suffix == ".pdf" or data.startswith(b"%PDF"):
        document = _extract_pdf(data)
    elif suffix == ".docx":
        document = _extract_docx(data)
    else:
        raise ResumeExtractionError("only PDF and DOCX files are supported")
    if len(document.text.strip()) < 100:
        raise ResumeExtractionError("no readable text found (is this a scanned image?)")
    return document


def detect_email(text: str) -> str | None:
    match = _EMAIL_RE.search(text)
    return match.group(0).lower() if match else None


def detect_phone(text: str) -> str | None:
    for match in _PHONE_RE.finditer(text):
        candidate = " ".join(match.group(0).split())
        if sum(ch.isdigit() for ch in candidate) >= 10 and not re.fullmatch(
            r"\d{4}\s*[-\u2013]\s*\d{4}", candidate
        ):
            return candidate
    return None


def _normalise_url(url: str) -> str | None:
    url = url.strip().rstrip(".")
    if url.lower().startswith(("mailto:", "tel:")) or "@" in url:
        return None
    full = url if re.match(r"^https?://", url, re.IGNORECASE) else f"https://{url}"
    host = (urlparse(full).hostname or "").lower()
    if "." not in host:
        return None
    return full


def find_urls(document: ExtractedDocument) -> list[str]:
    # Scan the redacted text so the domain half of an email address is never mistaken for a link.
    in_text = [c for c in _URL_RE.findall(redact_pii(document.text)) if _looks_like_url(c)]
    candidates = [*document.link_targets, *in_text]
    seen: dict[str, str] = {}
    for candidate in candidates:
        if (url := _normalise_url(candidate)) is not None:
            seen.setdefault(display_url(url).lower(), url)
    return list(seen.values())


def profile_links(urls: list[str]) -> list[Link]:
    links = []
    for url in urls:
        shown = display_url(url).lower()
        for label, pattern in _PROFILE_PATTERNS:
            if pattern.match(shown):
                links.append(Link(label=label, url=url))
                break
    return links


def _merge_links(found: list[Link], from_llm: list[Link], project_links: set[str]) -> list[Link]:
    merged: dict[str, Link] = {}
    for link in [*found, *from_llm]:
        key = display_url(link.url).lower()
        if key and key not in project_links and _normalise_url(link.url):
            merged.setdefault(key, link)
    return list(merged.values())


def _bullets(prefix: str, parsed: list[ParsedBullet], lexicon: SkillLexicon) -> list[Bullet]:
    return [
        Bullet(id=f"{prefix}-b{i}", text=b.text.strip(), skills=lexicon.normalize_all(b.skills))
        for i, b in enumerate(parsed, start=1)
        if b.text.strip()
    ]


def build_master_resume(
    parsed: ParsedResume,
    *,
    email: str | None,
    phone: str | None,
    found_links: list[Link],
    lexicon: SkillLexicon,
) -> MasterResume:
    """Assign stable ids, normalise skill tags, and re-insert the contact details kept from the LLM."""
    project_links = {display_url(p.link).lower() for p in parsed.projects if p.link}
    skills: dict[str, list[str]] = {}
    for group in parsed.skills:
        items = [i.strip() for i in group.items if i.strip()]
        if items and group.category.strip():
            existing = skills.setdefault(group.category.strip(), [])
            existing += [i for i in items if i not in existing]

    return MasterResume(
        basics=Basics(
            name=parsed.basics.name.strip(),
            email=email,
            phone=phone,
            location=parsed.basics.location,
            links=_merge_links(found_links, parsed.basics.links, project_links),
        ),
        summary=(parsed.summary or "").strip() or None,
        education=[
            Education(
                id=f"edu{i}",
                **e.model_dump(exclude={"bullets"}),
                bullets=_bullets(f"edu{i}", e.bullets, lexicon),
            )
            for i, e in enumerate(parsed.education, start=1)
        ],
        experience=[
            Experience(
                id=f"exp{i}",
                **e.model_dump(exclude={"bullets"}),
                bullets=_bullets(f"exp{i}", e.bullets, lexicon),
            )
            for i, e in enumerate(parsed.experience, start=1)
        ],
        projects=[
            Project(
                id=f"proj{i}",
                **p.model_dump(exclude={"bullets"}),
                bullets=_bullets(f"proj{i}", p.bullets, lexicon),
            )
            for i, p in enumerate(parsed.projects, start=1)
        ],
        skills=skills,
        achievements=_bullets("ach", parsed.achievements, lexicon),
    )


def build_parse_prompt(document: ExtractedDocument, urls: list[str]) -> tuple[str, str]:
    """(redacted resume text, hyperlink list) exactly as sent to the LLM."""
    text = redact_pii(document.text)[:MAX_TEXT_CHARS]
    links = "\n".join(f"- {u}" for u in urls) or "(none)"
    return text, links


async def parse_resume(data: bytes, filename: str, llm: LLMClient, api_key: str | None) -> MasterResume:
    document = await asyncio.to_thread(extract_document, data, filename)
    urls = find_urls(document)
    resume_text, links = build_parse_prompt(document, urls)
    request = load_prompt("parse_resume").request(
        tier="fast", api_key=api_key, resume_text=resume_text, links=links
    )
    parsed = await llm.generate(request, ParsedResume)
    return build_master_resume(
        parsed,
        email=detect_email(document.text),
        phone=detect_phone(document.text),
        found_links=profile_links(urls),
        lexicon=get_lexicon(),
    )
